import os
from datetime import date
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from config.settings import MIN_SCORE
from storage.repository import get_new_jobs, get_pipeline_jobs
from utils.logger import get_logger

logger = get_logger(__name__)

EXPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "exports")

GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
YELLOW = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def export_daily(run_date: date | None = None) -> str:
    run_date = run_date or date.today()
    os.makedirs(EXPORTS_DIR, exist_ok=True)
    filepath = os.path.join(EXPORTS_DIR, f"jobs_{run_date.isoformat()}.xlsx")

    new_jobs = get_new_jobs(run_date)
    pipeline = get_pipeline_jobs()

    top_jobs = pd.DataFrame()
    all_new = pd.DataFrame()

    if not new_jobs.empty:
        top_jobs = new_jobs[new_jobs["score"] >= MIN_SCORE].sort_values("score", ascending=False)
        all_new = new_jobs.sort_values("score", ascending=False)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        _write_new_sheet(writer, top_jobs, run_date)
        _write_all_sheet(writer, all_new)
        _write_track_sheet(writer, pipeline)

    _apply_formatting(filepath)
    logger.info(f"Excel exported: {filepath}")
    return filepath


def _new_cols() -> list[str]:
    return ["score", "source", "title", "company", "location", "is_remote",
            "salary", "date_posted", "stack_detected", "job_url", "description_short"]


def _prep_new(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=_new_cols())
    out = df.copy()
    out["salary"] = out.apply(
        lambda r: f"{r.get('min_amount','')}-{r.get('max_amount','')} {r.get('currency','')}"
        if r.get("min_amount") else "", axis=1
    )
    out["stack_detected"] = out["stack_detected"].apply(
        lambda x: ", ".join(x) if isinstance(x, list) else str(x)
    )
    out["description_short"] = out["description"].apply(
        lambda d: str(d)[:300] if d else ""
    )
    return out[_new_cols()]


def _write_new_sheet(writer, df: pd.DataFrame, run_date: date):
    prepped = _prep_new(df)
    prepped.to_excel(writer, sheet_name="NEW", index=False)


def _write_all_sheet(writer, df: pd.DataFrame):
    prepped = _prep_new(df)
    prepped.to_excel(writer, sheet_name="ALL", index=False)


def _track_cols() -> list[str]:
    return ["applied_at", "company", "title", "job_url", "status",
            "cv_version", "notes", "next_action_at", "response_at"]


def _write_track_sheet(writer, df: pd.DataFrame):
    if df.empty or "status" not in df.columns:
        empty = pd.DataFrame(columns=_track_cols())
        empty.to_excel(writer, sheet_name="TRACK", index=False)
        return

    applied = df[df["status"].notna() & (df["status"] != "Pendiente")].copy()
    if applied.empty:
        empty = pd.DataFrame(columns=_track_cols())
        empty.to_excel(writer, sheet_name="TRACK", index=False)
        return

    applied = applied[_track_cols()]
    applied.to_excel(writer, sheet_name="TRACK", index=False)


def _apply_formatting(filepath: str):
    wb = load_workbook(filepath)

    for sheet_name in ["NEW", "ALL"]:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        _style_header(ws)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # Color score column (col A = index 1)
        for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
            for cell in row:
                try:
                    val = int(cell.value or 0)
                    if val >= 70:
                        cell.fill = GREEN
                    elif val >= 50:
                        cell.fill = YELLOW
                    else:
                        cell.fill = RED
                except (TypeError, ValueError):
                    pass

        _autofit_columns(ws)

    if "TRACK" in wb.sheetnames:
        ws = wb["TRACK"]
        _style_header(ws)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        _add_status_dropdown(ws)
        _autofit_columns(ws)

    wb.save(filepath)


def _style_header(ws):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autofit_columns(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 60)


def _add_status_dropdown(ws):
    statuses = "Pendiente,Aplicado,Entrevista 1,Entrevista Téc,Oferta,Rechazado,Ghosted"
    dv = DataValidation(type="list", formula1=f'"{statuses}"', allow_blank=True)
    ws.add_data_validation(dv)
    # Status is column 5 (E) in TRACK sheet
    for row in range(2, ws.max_row + 1):
        dv.add(ws[f"E{row}"])
